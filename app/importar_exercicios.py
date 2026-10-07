"""Importa o quadro de exercícios (planilha) para o banco do app.

Uso:
    python -m app.importar_exercicios dados/quadro_exercicios_v2.xlsx app.db

Pode rodar mais de uma vez: o exercício que já existe (mesmo nome) é atualizado,
não duplicado. Exercícios criados dentro do app não são mexidos.
"""

import sqlite3
import sys
from pathlib import Path

import openpyxl

from app.db import conectar, criar_tabelas

# Os 15 grupos, na ordem em que aparecem na tela de montar treino.
GRUPOS = [
    "Peito", "Costas", "Ombros", "Bíceps", "Tríceps", "Antebraços",
    "Quadríceps", "Posterior de coxa", "Glúteos", "Quadril (adutores/abdutores)",
    "Panturrilhas", "Abdômen", "Lombar", "Cardio", "Mobilidade",
]

# Como a coluna "Origem" da planilha vira o valor guardado no banco.
ORIGENS = {
    "criado na academia": "data4u_academia",
    "catalogo Data4U (com animacao)": "data4u_catalogo",
}

COLUNAS = {
    "nome": "Exercício",
    "grupo_sugerido": "Grupo sugerido",
    "grupo_final": "Grupo final (corrigir aqui)",
    "confianca": "Confiança",
    "usos": "Prescrições no último ano",
    "ids": "IDs Data4U (todos)",
    "origem": "Origem",
}

CONFIANCA_COMBINADO = "Revisar (combinado)"


class ErroDeImportacao(ValueError):
    """A planilha tem algo que não dá para importar com segurança."""


def _ler_quadro(caminho: str | Path) -> list[dict]:
    """Lê a aba 'Quadro' e devolve uma lista de dicionários já validados."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        if "Quadro" not in wb.sheetnames:
            raise ErroDeImportacao("A planilha não tem a aba 'Quadro'.")
        linhas = list(wb["Quadro"].iter_rows(values_only=True))
    finally:
        wb.close()

    if not linhas:
        raise ErroDeImportacao("A aba 'Quadro' está vazia.")
    cabecalho = [str(c).strip() if c is not None else "" for c in linhas[0]]
    faltando = [c for c in COLUNAS.values() if c not in cabecalho]
    if faltando:
        raise ErroDeImportacao(f"Faltam colunas na aba 'Quadro': {', '.join(faltando)}")
    pos = {chave: cabecalho.index(titulo) for chave, titulo in COLUNAS.items()}

    exercicios: list[dict] = []
    nomes_vistos: set[str] = set()
    ids_vistos: dict[int, str] = {}
    for numero, linha in enumerate(linhas[1:], start=2):  # 2 = primeira linha de dados
        if all(c is None for c in linha):
            continue  # linha totalmente vazia no fim da planilha

        nome = str(linha[pos["nome"]] or "").strip()
        if not nome:
            raise ErroDeImportacao(f"Linha {numero}: exercício sem nome.")
        if nome.casefold() in nomes_vistos:
            raise ErroDeImportacao(f"Linha {numero}: nome repetido: {nome}")
        nomes_vistos.add(nome.casefold())

        # O grupo corrigido à mão vale mais que o sugerido.
        texto_grupo = linha[pos["grupo_final"]] or linha[pos["grupo_sugerido"]]
        grupos = [g.strip() for g in str(texto_grupo).split(" + ")] if texto_grupo else []
        desconhecidos = [g for g in grupos if g not in GRUPOS]
        if desconhecidos:
            raise ErroDeImportacao(f"Linha {numero} ({nome}): grupo desconhecido: {desconhecidos[0]}")

        ids = []
        for parte in str(linha[pos["ids"]] or "").split(","):
            parte = parte.strip()
            if not parte:
                continue
            try:
                id_data4u = int(parte)
            except ValueError:
                raise ErroDeImportacao(f"Linha {numero} ({nome}): id do Data4U inválido: {parte!r}") from None
            if id_data4u in ids_vistos:
                raise ErroDeImportacao(
                    f"Linha {numero} ({nome}): id {id_data4u} já está em '{ids_vistos[id_data4u]}'"
                )
            ids_vistos[id_data4u] = nome
            ids.append(id_data4u)
        if not ids:
            raise ErroDeImportacao(f"Linha {numero} ({nome}): sem id do Data4U.")

        origem_texto = str(linha[pos["origem"]] or "").strip()
        if origem_texto not in ORIGENS:
            raise ErroDeImportacao(f"Linha {numero} ({nome}): origem desconhecida: {origem_texto!r}")

        confianca = str(linha[pos["confianca"]]).strip() if linha[pos["confianca"]] else None
        exercicios.append({
            "nome": nome,
            "grupos": grupos,
            "confianca": confianca,
            "combinado": 1 if confianca == CONFIANCA_COMBINADO else 0,
            "usos": int(linha[pos["usos"]] or 0),
            "ids": ids,
            "origem": ORIGENS[origem_texto],
        })
    return exercicios


def importar_quadro(conn: sqlite3.Connection, caminho: str | Path) -> dict:
    """Grava o quadro no banco, tudo ou nada, e devolve um resumo com as contagens."""
    exercicios = _ler_quadro(caminho)  # se algo estiver errado, falha aqui, antes de gravar

    with conn:  # transação: se der erro no meio, nada fica pela metade
        for ordem, nome in enumerate(GRUPOS, start=1):
            conn.execute(
                "INSERT INTO grupo_muscular (nome, ordem) VALUES (?, ?) "
                "ON CONFLICT(nome) DO UPDATE SET ordem = excluded.ordem",
                (nome, ordem),
            )
        id_do_grupo = {r["nome"]: r["id"] for r in conn.execute("SELECT id, nome FROM grupo_muscular")}

        for ex in exercicios:
            conn.execute(
                "INSERT INTO exercicio (nome, combinado, origem, usos_ultimo_ano, confianca_grupo) "
                "VALUES (?, ?, ?, ?, ?) "
                "ON CONFLICT(nome) DO UPDATE SET combinado = excluded.combinado, "
                "origem = excluded.origem, usos_ultimo_ano = excluded.usos_ultimo_ano, "
                "confianca_grupo = excluded.confianca_grupo",
                (ex["nome"], ex["combinado"], ex["origem"], ex["usos"], ex["confianca"]),
            )
            exercicio_id = conn.execute("SELECT id FROM exercicio WHERE nome = ?", (ex["nome"],)).fetchone()["id"]

            conn.execute("DELETE FROM exercicio_grupo WHERE exercicio_id = ?", (exercicio_id,))
            for grupo in ex["grupos"]:
                conn.execute(
                    "INSERT INTO exercicio_grupo (exercicio_id, grupo_id) VALUES (?, ?)",
                    (exercicio_id, id_do_grupo[grupo]),
                )
            for id_data4u in ex["ids"]:
                conn.execute(
                    "INSERT INTO exercicio_data4u (data4u_id, exercicio_id) VALUES (?, ?) "
                    "ON CONFLICT(data4u_id) DO UPDATE SET exercicio_id = excluded.exercicio_id",
                    (id_data4u, exercicio_id),
                )

    return {
        "exercicios": len(exercicios),
        "combinados": sum(e["combinado"] for e in exercicios),
        "sem_grupo": sum(1 for e in exercicios if not e["grupos"]),
        "ids_data4u": sum(len(e["ids"]) for e in exercicios),
    }


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("Uso: python -m app.importar_exercicios <quadro.xlsx> <banco.db>")
        return 2
    conn = conectar(argv[2])
    criar_tabelas(conn)
    try:
        resumo = importar_quadro(conn, argv[1])
    except ErroDeImportacao as erro:
        print(f"Importação cancelada, nada foi gravado: {erro}")
        return 1
    print("Importação concluída:", resumo)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
